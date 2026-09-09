import io

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook

from inspection_nlp.documents import NO_OCR, parse_document


def test_pdf_text_and_page_provenance():
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "A valve cracked.")
        data = document.tobytes()
    records = parse_document(data, "report.pdf")
    assert records[0].clean_text == "A valve cracked."
    assert records[0].document_metadata["page"] == "1"


def test_image_only_pdf_rejected():
    with pymupdf.open() as document:
        document.new_page()
        data = document.tobytes()
    with pytest.raises(ValueError, match="Scanned documents") as error:
        parse_document(data, "scan.pdf")
    assert NO_OCR in str(error.value)


def test_docx_paragraph_table_order():
    document = Document()
    document.add_paragraph("Inspection start.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Valve"
    table.cell(0, 1).text = "cracked"
    document.add_paragraph("Repair required.")
    buffer = io.BytesIO()
    document.save(buffer)
    text = parse_document(buffer.getvalue(), "report.docx")[0].clean_text
    assert text == "Inspection start. Valve | cracked Repair required."


@pytest.mark.parametrize(
    "name,delimiter,tabular",
    [("report.csv", ",", False), ("report.tsv", "\t", False), ("report.txt", "\t", True)],
)
def test_tables_select_only_narrative(name, delimiter, tabular):
    data = f"Narrative{delimiter}Target\nValve inspected.{delimiter}CRACKED\n".encode()
    records = parse_document(data, name, text_column="Narrative", tabular_txt=tabular)
    assert records[0].clean_text == "Valve inspected."
    assert "CRACKED" not in records[0].model_dump_json()
    assert records[0].source_row == 2


def test_xlsx_and_empty_rows():
    workbook = Workbook()
    workbook.active.append(["Narrative", "Target"])
    workbook.active.append([None, "leak"])
    workbook.active.append(["Pipe inspected.", "leak"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    records = parse_document(buffer.getvalue(), "report.xlsx", text_column="Narrative")
    assert len(records) == 1
    assert records[0].source_row == 3
    assert records[0].clean_text == "Pipe inspected."


def test_missing_column_and_bad_input():
    with pytest.raises(ValueError, match="Select the narrative"):
        parse_document(b"text\ncracked", "report.csv")
    with pytest.raises(ValueError, match="exactly once"):
        parse_document(b"text\ncracked", "report.csv", text_column="wrong")
    with pytest.raises(ValueError, match="No narrative"):
        parse_document(b"  ", "empty.txt")
    with pytest.raises(ValueError, match="Supported formats"):
        parse_document(b"abc", "report.exe")


def test_text_redaction():
    record = parse_document(b"Email a@example.com about the leak.", "note.txt")[0]
    assert record.contains_sensitive_fields
    assert "a@example.com" not in record.display_text
