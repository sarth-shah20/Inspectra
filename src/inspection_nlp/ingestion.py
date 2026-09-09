"""Streaming source adapters. Only the declared narrative becomes model input."""

import csv
import hashlib
from collections.abc import Iterator
from pathlib import Path

from openpyxl import load_workbook

from .preprocessing import prepare_text
from .schemas import Record


def checksum(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def rows(path: Path) -> Iterator[tuple[int, dict[str, str]]]:
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            values = workbook.active.iter_rows(values_only=True)
            headers = list(next(values))
            for index, values_row in enumerate(values, 2):
                yield (
                    index,
                    dict(
                        zip(
                            headers,
                            [str(v) if v is not None else "" for v in values_row],
                            strict=True,
                        )
                    ),
                )
        finally:
            workbook.close()
    else:
        # Strict decoding: never silently corrupt source evidence.
        encoding = "cp1252" if path.suffix == ".txt" else "utf-8-sig"
        with path.open(encoding=encoding, newline="") as stream:
            reader = csv.DictReader(stream, delimiter="\t" if path.suffix == ".txt" else ",")
            for index, row in enumerate(reader, 2):
                if None in row or any(v is None for v in row.values()):
                    raise ValueError(f"Malformed row in {path.name}:{index}")
                yield index, row


def load_source(path: Path) -> Iterator[Record]:
    digest = checksum(path)
    for index, row in rows(path):
        if path.suffix == ".xlsx":
            source, domain, version = "fire_door", "construction", "fire_door_v1"
            text_col, id_col, date_col = "English", "", ""
            required = {
                "Apartments",
                "English",
                "Classification_English",
            }
            version = "fire_door_bilingual_v1" if "Korean" in row else "fire_door_english_v1"
            targets = ["Classification_English"]
        elif path.suffix == ".csv":
            source, domain, version = "faa_sdr", "aviation", "faa_sdr_v1"
            text_col, id_col, date_col = "Discrepancy", "OperatorControlNumber", "DifficultyDate"
            required = {text_col, id_col}
            targets = ["PartCondition", "PartName", "ComponentName"]
        else:
            source, domain = "phmsa", "pipeline"
            legacy = "mar2004_dec2009" in path.stem
            version = "phmsa_legacy_v1" if legacy else "phmsa_current_v1"
            text_col = "NARRATIVE"
            id_col = "RPTID" if legacy else "REPORT_NUMBER"
            date_col = "IDATE" if legacy else "LOCAL_DATETIME"
            required = {text_col, id_col}
            targets = (
                ["CAUSE", "CAUSE_TEXT", "CAUSE_DETAILS_TEXT"]
                if legacy
                else ["CAUSE", "CAUSE_DETAILS", "SYSTEM_PART_INVOLVED", "MATERIAL_INVOLVED"]
            )
        if missing := required - row.keys():
            raise ValueError(f"{path.name}: missing required columns {sorted(missing)}")
        original_id = row[id_col].strip() if id_col else f"{path.stem}:{index}"
        if not original_id:
            raise ValueError(f"{path.name}:{index}: missing event identifier")
        text, sensitive = prepare_text(row[text_col])
        metadata = {
            "source_subset": path.stem,
            "privacy_review": "pending",
            "encoding": "cp1252" if path.suffix == ".txt" else "utf-8-sig",
        }
        if source == "fire_door":
            metadata.update(original_split=path.stem, translation="source_provided_English")
        if not text:
            metadata["exclusion_reason"] = "empty_narrative"
        yield Record(
            record_id=f"{source}:{digest[:16]}:{index}",
            source_dataset=source,
            source_record_id=original_id,
            source_event_id=f"{source}:{original_id}",
            source_document=path.name,
            source_sha256=digest,
            source_row=index,
            domain=domain,
            schema_version=version,
            source_column_mapping={
                "clean_text": text_col,
                "source_record_id": id_col or "row",
                "report_date": date_col,
            },
            report_date=row.get(date_col) or None,
            report_status=row.get("REPORT_TYPE") or None,
            clean_text=text,
            display_text=text,
            contains_sensitive_fields=sensitive,
            structured_source_fields={k: prepare_text(row[k])[0] for k in targets if row.get(k)},
            document_metadata=metadata,
        )
