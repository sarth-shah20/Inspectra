"""User-document adapters: one record per page, document, or selected table row."""

import csv
import hashlib
import io
from pathlib import Path

from .metadata import FIELDS, normalize_metadata, vendor_key
from .preprocessing import prepare_evidence
from .schemas import Domain, Record

NO_OCR = (
    "This version supports machine-readable PDFs only. "
    "Scanned documents require OCR and are outside the project scope."
)
MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 10000


def _record(
    text: str,
    name: str,
    digest: str,
    index: int,
    domain: Domain,
    text_column: str = "text",
    metadata: dict | None = None,
    report_metadata: dict | None = None,
    provenance: dict | None = None,
    aliases: dict | None = None,
    table: bool = False,
) -> Record:
    display, cleaned, offsets, sensitive = prepare_evidence(text)
    values = dict(report_metadata or {})
    if values.get("vendor"):
        values["vendor"] = vendor_key(values["vendor"], aliases)
    group = values.get("report_number") or (str(index) if table else "document")
    return Record(
        record_id=f"upload:{digest}:{index}",
        source_dataset="user_upload",
        source_record_id=str(index),
        source_event_id=f"upload:{digest}:{index}",
        source_document=Path(name).name,
        source_sha256=digest,
        source_row=index + 2,
        domain=domain,
        schema_version="upload_v2",
        source_column_mapping={"clean_text": text_column},
        clean_text=cleaned,
        display_text=display,
        clean_to_display=offsets,
        report_id=f"upload:{digest}:{group}",
        report_metadata=values,
        metadata_provenance=provenance or {},
        report_date=values.get("report_date"),
        quality_flags=["sparse_text"] if len(cleaned.split()) < 4 else [],
        contains_sensitive_fields=sensitive,
        document_metadata={"privacy_review": "pending", **(metadata or {})},
    )


def parse_document(
    data: bytes,
    filename: str,
    *,
    domain: Domain = "general",
    text_column: str | None = None,
    encoding: str = "utf-8-sig",
    tabular_txt: bool = False,
    metadata_columns: dict[str, str] | None = None,
    report_metadata: dict[str, str] | None = None,
    date_format: str | None = None,
    vendor_aliases: dict[str, str] | None = None,
) -> list[Record]:
    if len(data) > MAX_BYTES:
        raise ValueError("Upload exceeds the 20 MiB limit.")
    metadata_columns = metadata_columns or {}
    if set(metadata_columns) - set(FIELDS):
        raise ValueError("Unknown metadata column mapping")
    supplied = normalize_metadata(report_metadata or {}, date_format)
    row_values = {}
    table = False
    digest = hashlib.sha256(data).hexdigest()
    suffix = Path(filename).suffix.lower()
    segments: list[tuple[str, int, dict]] = []
    if suffix == ".pdf":
        import pymupdf

        with pymupdf.open(stream=data, filetype="pdf") as document:
            if document.needs_pass:
                raise ValueError("Password-protected PDFs are not supported.")
            for index, page in enumerate(document):
                text = page.get_text(sort=True)
                if not text.strip():
                    # Reject mixed scanned/text files rather than silently omit pages.
                    raise ValueError(f"Page {index + 1}: {NO_OCR}")
                segments.append((text, index, {"page": str(index + 1)}))
    elif suffix == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(io.BytesIO(data))
        blocks = []
        for item in document.iter_inner_content():
            if isinstance(item, Paragraph):
                blocks.append(item.text)
            elif isinstance(item, Table):
                blocks.extend(" | ".join(cell.text for cell in row.cells) for row in item.rows)
        segments.append(("\n".join(blocks), 0, {}))
    elif suffix in {".csv", ".tsv", ".xlsx"} or (suffix == ".txt" and tabular_txt):
        table = True
        if not text_column:
            raise ValueError("Select the narrative text column; other columns are not model input.")
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                values = workbook.active.iter_rows(values_only=True)
                headers = list(next(values, ()))
                if headers.count(text_column) != 1:
                    raise ValueError("Narrative column must exist exactly once.")
                for mapped in metadata_columns.values():
                    if headers.count(mapped) != 1:
                        raise ValueError(f"Metadata column {mapped} must exist exactly once")
                column = headers.index(text_column)
                for index, row in enumerate(values):
                    if index >= MAX_ROWS:
                        raise ValueError("Table exceeds the 10,000-row upload limit.")
                    row_values[index] = normalize_metadata(
                        {
                            field: row[headers.index(mapped)]
                            for field, mapped in metadata_columns.items()
                        },
                        date_format,
                    )
                    value = row[column]
                    segments.append((str(value) if value is not None else "", index, {}))
            finally:
                workbook.close()
        else:
            stream = io.StringIO(data.decode(encoding), newline="")
            reader = csv.DictReader(stream, delimiter="," if suffix == ".csv" else "\t")
            if not reader.fieldnames or reader.fieldnames.count(text_column) != 1:
                raise ValueError("Narrative column must exist exactly once.")
            for mapped in metadata_columns.values():
                if reader.fieldnames.count(mapped) != 1:
                    raise ValueError(f"Metadata column {mapped} must exist exactly once")
            for index, row in enumerate(reader):
                if index >= MAX_ROWS:
                    raise ValueError("Table exceeds the 10,000-row upload limit.")
                if None in row or any(v is None for v in row.values()):
                    raise ValueError(f"Malformed table row {index + 2}.")
                row_values[index] = normalize_metadata(
                    {field: row[mapped] for field, mapped in metadata_columns.items()}, date_format
                )
                segments.append((row[text_column], index, {}))
    elif suffix == ".txt":
        segments.append((data.decode(encoding), 0, {}))
    else:
        raise ValueError("Supported formats: PDF, DOCX, CSV, XLSX, TSV, and TXT.")
    records = [
        _record(
            text,
            filename,
            digest,
            index,
            domain,
            text_column or "text",
            metadata,
            {**supplied, **row_values.get(index, {})},
            {
                **{key: "user" for key in supplied},
                **{key: f"column:{metadata_columns[key]}" for key in row_values.get(index, {})},
            },
            vendor_aliases,
            table,
        )
        for text, index, metadata in segments
        if text.strip()
    ]
    empty = sum(not text.strip() for text, _, _ in segments)
    if empty:
        for record in records:
            record.quality_flags.append(f"empty_rows:{empty}")
    if not records:
        raise ValueError("No narrative text found in the selected input.")
    return records


def table_headers(data: bytes, filename: str, encoding: str = "utf-8-sig") -> list[str]:
    if Path(filename).suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            return [str(x) for x in next(workbook.active.values, ()) if x is not None]
        finally:
            workbook.close()
    reader = csv.reader(
        io.StringIO(data.decode(encoding)),
        delimiter="," if filename.lower().endswith(".csv") else "\t",
    )
    return next(reader, [])


def parse_batch(files: list[tuple[str, bytes]], **settings) -> tuple[list[Record], list[dict]]:
    from zipfile import BadZipFile

    records, errors = [], []
    for name, data in files:
        try:
            records.extend(parse_document(data, name, **settings))
        except (ValueError, UnicodeError, RuntimeError, BadZipFile, KeyError, IndexError) as exc:
            errors.append(
                {"file": name, "error": str(exc), "quality_flag": "unsupported_or_malformed"}
            )
    return records, errors
